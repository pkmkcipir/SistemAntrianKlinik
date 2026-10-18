"""Laporan mingguan/bulanan: keuangan dan obat, plus ekspor ke Excel (.xlsx) dan PDF."""
import re
from datetime import date, datetime, timedelta
from html import escape

from .config import BULAN, COPYRIGHT, HARI, angka, fmt_date, month_range, rupiah, week_range


def kop(db):
    """Kop laporan: identitas klinik dari menu Pengaturan."""
    return {"nama": db.get("nama_klinik") or "", "alamat": db.get("alamat") or "",
            "telepon": db.get("telepon") or ""}


def format_cell(v, jenis):
    if v is None:
        return ""
    if jenis == "uang":
        return rupiah(v)
    if jenis == "angka":
        return angka(v)
    return str(v)


def laporan(db, jenis, periode, anchor: date):
    """jenis: 'Keuangan' | 'Obat'; periode: 'Mingguan' | 'Bulanan'; anchor: tanggal acuan."""
    if periode == "Mingguan":
        s, e = week_range(anchor)
        label = f"Minggu {fmt_date(s)} s/d {fmt_date(e)}"
    else:
        s, e = month_range(anchor)
        label = f"Bulan {BULAN[anchor.month - 1]} {anchor.year}"
    rep = _keuangan(db, s, e) if jenis == "Keuangan" else _obat(db, s, e)
    rep.update(judul=f"Laporan {jenis} {periode}", periode=label, mulai=s, selesai=e,
               nama_file=f"Laporan-{jenis}-{periode}-{s.isoformat()}")
    return rep


def _keuangan(db, s, e):
    rows = db.keu_list(s.isoformat(), e.isoformat())
    masuk = sum(r["jumlah"] for r in rows if r["tipe"] == "Pemasukan")
    keluar = sum(r["jumlah"] for r in rows if r["tipe"] == "Pengeluaran")

    harian, d = {}, s
    while d <= e:
        harian[d.isoformat()] = [0, 0]
        d += timedelta(days=1)
    for r in rows:
        harian.setdefault(r["tanggal"], [0, 0])[0 if r["tipe"] == "Pemasukan" else 1] += r["jumlah"]

    t1 = {"judul": "Rekap Harian", "total": True,
          "kolom": [("Tanggal", "teks"), ("Pemasukan", "uang"), ("Pengeluaran", "uang"), ("Saldo", "uang")],
          "baris": [[f"{HARI[date.fromisoformat(k).weekday()]}, {fmt_date(k)}", v[0], v[1], v[0] - v[1]]
                    for k, v in sorted(harian.items())] + [["TOTAL", masuk, keluar, masuk - keluar]]}

    kat = {}
    for r in rows:
        x = kat.setdefault((r["tipe"], r["kategori"] or "(tanpa kategori)"), [0, 0])
        x[0] += 1
        x[1] += r["jumlah"]
    t2 = {"judul": "Per Kategori",
          "kolom": [("Tipe", "teks"), ("Kategori", "teks"), ("Jumlah Transaksi", "angka"), ("Total", "uang")],
          "baris": [[k[0], k[1], v[0], v[1]] for k, v in
                    sorted(kat.items(), key=lambda kv: (kv[0][0] != "Pemasukan", -kv[1][1]))]}

    t3 = {"judul": "Rincian Transaksi",
          "kolom": [("Tanggal", "teks"), ("Tipe", "teks"), ("Kategori", "teks"), ("Keterangan", "teks"),
                    ("Pasien", "teks"), ("Metode", "teks"), ("Jumlah", "uang")],
          "baris": [[fmt_date(r["tanggal"]), r["tipe"], r["kategori"], r["keterangan"],
                     r["pasien_nama"] or "", r["metode"], r["jumlah"]] for r in rows]}

    return {"ringkasan": [("Total Pemasukan", rupiah(masuk)), ("Total Pengeluaran", rupiah(keluar)),
                          ("Saldo", rupiah(masuk - keluar)), ("Jumlah Transaksi", angka(len(rows)))],
            "tabel": [t1, t2, t3]}


def _obat(db, s, e):
    baris, tot_masuk, tot_keluar, menipis = [], 0, 0, 0
    for o in db.obat_rekap(s.isoformat(), e.isoformat()):
        gerak = o["awal"] or o["masuk"] or o["keluar"] or o["koreksi"]
        if not o["aktif"] and not gerak:
            continue
        akhir = o["awal"] + o["masuk"] - o["keluar"] + o["koreksi"]
        status = "Habis" if akhir <= 0 else ("Menipis" if akhir <= o["stok_min"] else "Aman")
        menipis += status != "Aman"
        tot_masuk += o["masuk"]
        tot_keluar += o["keluar"]
        baris.append([o["kode"], o["nama"], o["satuan"], o["awal"], o["masuk"], o["keluar"],
                      o["koreksi"], akhir, o["stok_min"], status])
    t1 = {"judul": "Rekap Stok Obat",
          "kolom": [("Kode", "teks"), ("Nama Obat", "teks"), ("Satuan", "teks"), ("Stok Awal", "angka"),
                    ("Masuk", "angka"), ("Keluar", "angka"), ("Koreksi", "angka"), ("Stok Akhir", "angka"),
                    ("Stok Min", "angka"), ("Status", "teks")],
          "baris": baris}
    t2 = {"judul": "Rincian Mutasi",
          "kolom": [("Tanggal", "teks"), ("Obat", "teks"), ("Tipe", "teks"), ("Jumlah", "angka"),
                    ("Satuan", "teks"), ("Keterangan", "teks"), ("Pasien", "teks")],
          "baris": [[fmt_date(m["tanggal"]), m["obat_nama"], m["tipe"], m["delta"], m["satuan"],
                     m["keterangan"], m["pasien_nama"] or ""]
                    for m in db.mutasi_list(s.isoformat(), e.isoformat())]}
    return {"ringkasan": [("Jenis Obat", angka(len(baris))), ("Total Masuk", angka(tot_masuk)),
                          ("Total Keluar", angka(tot_keluar)), ("Menipis / Habis", angka(menipis))],
            "tabel": [t1, t2]}


# ---- ekspor ------------------------------------------------------------------
def export_xlsx(rep, path, k):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for t in rep["tabel"]:
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "-", t["judul"])[:31])
        ws.append([k["nama"]])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([k["alamat"]])
        ws.append([f"Telp/WA {k['telepon']}" if k["telepon"] else ""])
        ws.append([])
        ws.append([rep["judul"]])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append([rep["periode"]])
        for lab, val in rep["ringkasan"]:
            ws.append([lab, val])
        ws.append([])
        head = ws.max_row + 1
        ws.append([n for n, _ in t["kolom"]])
        for cell in ws[head]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0E6B67")
            cell.alignment = Alignment(horizontal="center")
        for i, row in enumerate(t["baris"]):
            ws.append(row)
            for ci, (_n, jenis) in enumerate(t["kolom"], start=1):
                cell = ws.cell(ws.max_row, ci)
                if jenis == "uang":
                    cell.number_format = '"Rp" #,##0'
                elif jenis == "angka":
                    cell.number_format = "#,##0"
                if t.get("total") and i == len(t["baris"]) - 1:
                    cell.font = Font(bold=True)
        ws.append([])
        ws.append([COPYRIGHT])
        for ci, (n, jenis) in enumerate(t["kolom"], start=1):
            lebar = max([len(n)] + [len(format_cell(r[ci - 1], jenis)) for r in t["baris"]])
            ws.column_dimensions[get_column_letter(ci)].width = min(lebar + 3, 48)
        ws.freeze_panes = ws.cell(head + 1, 1)
    wb.save(path)


def _html(rep, k):
    h = ['<body style="font-family:Segoe UI,Arial,sans-serif; font-size:9pt; color:#10252b">']
    h.append(f'<div align="center"><b style="font-size:15pt">{escape(k["nama"])}</b><br>{escape(k["alamat"])}'
             + (f'<br>Telp/WA {escape(k["telepon"])}' if k["telepon"] else "") + "</div><hr>")
    h.append(f'<h3 style="margin-bottom:0">{escape(rep["judul"])}</h3>'
             f'<p style="margin-top:2px">{escape(rep["periode"])}</p>')
    h.append('<table cellspacing="0" cellpadding="3">' + "".join(
        f"<tr><td>{escape(a)}</td><td>: <b>{escape(b)}</b></td></tr>" for a, b in rep["ringkasan"]) + "</table>")
    for t in rep["tabel"]:
        h.append(f'<h4 style="margin-bottom:2px">{escape(t["judul"])}</h4>')
        h.append('<table width="100%" border="1" cellspacing="0" cellpadding="4">')
        h.append("<tr>" + "".join(f'<th bgcolor="#dfeeea">{escape(n)}</th>' for n, _ in t["kolom"]) + "</tr>")
        for i, row in enumerate(t["baris"]):
            tebal = t.get("total") and i == len(t["baris"]) - 1
            sel = ""
            for v, (_n, jenis) in zip(row, t["kolom"]):
                teks = escape(format_cell(v, jenis))
                sel += (f'<td align="{"right" if jenis in ("uang", "angka") else "left"}">'
                        f'{f"<b>{teks}</b>" if tebal else teks}</td>')
            h.append(f"<tr>{sel}</tr>")
        if not t["baris"]:
            h.append(f'<tr><td colspan="{len(t["kolom"])}" align="center">Tidak ada data pada periode ini</td></tr>')
        h.append("</table>")
    h.append(f'<p align="center" style="color:#5f7a76; font-size:8pt">Dicetak {datetime.now():%d-%m-%Y %H:%M}. '
             f'{escape(COPYRIGHT)}</p></body>')
    return "".join(h)


def export_pdf(rep, path, k):
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    doc = QTextDocument()
    doc.setHtml(_html(rep, k))
    w = QPdfWriter(path)
    w.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    w.setPageMargins(QMarginsF(14, 14, 14, 14), QPageLayout.Unit.Millimeter)
    doc.print_(w)
